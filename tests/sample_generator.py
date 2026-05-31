from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class ValidationSample:
    path: Path
    domain: str
    expected_header_row: int
    expected_data_start_row: int
    expected_mappings: dict[str, str]


VALUE_BY_TARGET = {
    "order_no": ["SO001", "SO002", "SO003"],
    "customer_name": ["腾讯", "阿里", "字节"],
    "amount": [1000, 2000.5, 3500],
    "order_date": ["2026-05-01", "2026-05-02", "2026-05-03"],
    "product_name": ["笔记本电脑", "显示器", "键盘"],
    "quantity": [2, 5, 8],
    "sku": ["SKU-001", "SKU-002", "SKU-003"],
    "inventory_qty": [120, 80, 45],
    "contact_name": ["张三", "李四", "王五"],
    "phone": ["13800138000", "13900139000", "13700137000"],
    "invoice_no": ["INV001", "INV002", "INV003"],
    "paid_status": ["是", "否", "是"],
    "account_name": ["基本户", "一般户", "收款户"],
    "department": ["销售部", "财务部", "运营部"],
    "internal_note": ["备注A", "备注B", "备注C"],
}


SPECS = [
    {
        "domain": "order",
        "name": "order_title_basic",
        "title": "销售订单导入模板",
        "header_row": 2,
        "fields": [
            ("订单号", "order_no"),
            ("购买方", "customer_name"),
            ("订单金额", "amount"),
            ("下单日期", "order_date"),
            ("商品名称", "product_name"),
            ("数量", "quantity"),
        ],
    },
    {
        "domain": "order",
        "name": "order_shifted_aliases",
        "title": "订单信息",
        "header_row": 3,
        "fields": [
            ("销售单号", "order_no"),
            ("企业名称", "customer_name"),
            ("含税金额", "amount"),
            ("业务日期", "order_date"),
            ("产品名称", "product_name"),
            ("订购数量", "quantity"),
        ],
    },
    {
        "domain": "order",
        "name": "order_reordered",
        "header_row": 1,
        "fields": [
            ("客户名称", "customer_name"),
            ("订单编号", "order_no"),
            ("总金额", "amount"),
            ("SKU", "sku"),
            ("销售数量", "quantity"),
        ],
    },
    {
        "domain": "order",
        "name": "order_deep_header",
        "title": "销售订单",
        "section": "订单信息",
        "header_row": 4,
        "fields": [
            ("单号", "order_no"),
            ("购买方", "customer_name"),
            ("应收金额", "amount"),
            ("日期", "order_date"),
            ("品名", "product_name"),
            ("采购数量", "quantity"),
        ],
    },
    {
        "domain": "order",
        "name": "order_hidden_column",
        "title": "订单导入",
        "header_row": 2,
        "hidden_columns": [7],
        "fields": [
            ("订单号", "order_no"),
            ("公司名称", "customer_name"),
            ("金额", "amount"),
            ("订单日期", "order_date"),
            ("商品编码", "sku"),
            ("数量", "quantity"),
            ("内部备注", "internal_note"),
        ],
    },
    {
        "domain": "crm",
        "name": "crm_basic",
        "title": "CRM客户导入",
        "header_row": 2,
        "fields": [
            ("客户名称", "customer_name"),
            ("联系人", "contact_name"),
            ("电话", "phone"),
            ("所属部门", "department"),
        ],
    },
    {
        "domain": "crm",
        "name": "crm_shifted",
        "title": "客户资料",
        "header_row": 3,
        "fields": [
            ("企业名称", "customer_name"),
            ("联系人姓名", "contact_name"),
            ("手机号", "phone"),
            ("业务部门", "department"),
        ],
    },
    {
        "domain": "crm",
        "name": "crm_reordered",
        "header_row": 1,
        "fields": [
            ("公司名称", "customer_name"),
            ("客户联系人", "contact_name"),
            ("联系电话", "phone"),
            ("部门", "department"),
        ],
    },
    {
        "domain": "crm",
        "name": "crm_contact_only",
        "title": "联系人清单",
        "header_row": 2,
        "fields": [
            ("客户名", "customer_name"),
            ("联系人", "contact_name"),
            ("手机", "phone"),
        ],
    },
    {
        "domain": "crm",
        "name": "crm_multi_sheet",
        "title": "客户联系表",
        "header_row": 2,
        "extra_sheet": True,
        "fields": [
            ("购买方", "customer_name"),
            ("联系人", "contact_name"),
            ("电话", "phone"),
            ("部门", "department"),
        ],
    },
    {
        "domain": "inventory",
        "name": "inventory_basic",
        "title": "库存模板",
        "header_row": 2,
        "fields": [
            ("SKU", "sku"),
            ("库存数量", "inventory_qty"),
            ("商品名称", "product_name"),
        ],
    },
    {
        "domain": "inventory",
        "name": "inventory_aliases",
        "header_row": 1,
        "fields": [
            ("商品编码", "sku"),
            ("库存", "inventory_qty"),
            ("产品名称", "product_name"),
        ],
    },
    {
        "domain": "inventory",
        "name": "inventory_material",
        "title": "物料库存",
        "header_row": 3,
        "fields": [
            ("物料编码", "sku"),
            ("现存量", "inventory_qty"),
            ("物料名称", "product_name"),
        ],
    },
    {
        "domain": "inventory",
        "name": "inventory_available",
        "header_row": 1,
        "fields": [
            ("产品编码", "sku"),
            ("可用库存", "inventory_qty"),
            ("品名", "product_name"),
        ],
    },
    {
        "domain": "inventory",
        "name": "inventory_department",
        "title": "仓库库存",
        "header_row": 3,
        "fields": [
            ("SKU", "sku"),
            ("库存数量", "inventory_qty"),
            ("商品名称", "product_name"),
            ("部门", "department"),
        ],
    },
    {
        "domain": "finance",
        "name": "finance_basic",
        "title": "财务发票",
        "header_row": 2,
        "fields": [
            ("发票号", "invoice_no"),
            ("金额", "amount"),
            ("付款状态", "paid_status"),
            ("账户名称", "account_name"),
        ],
    },
    {
        "domain": "finance",
        "name": "finance_aliases",
        "header_row": 1,
        "fields": [
            ("发票号码", "invoice_no"),
            ("含税金额", "amount"),
            ("是否付款", "paid_status"),
            ("开户名", "account_name"),
        ],
    },
    {
        "domain": "finance",
        "name": "finance_shifted",
        "title": "收款登记",
        "header_row": 3,
        "fields": [
            ("票据号", "invoice_no"),
            ("总金额", "amount"),
            ("是否支付", "paid_status"),
            ("银行账户名", "account_name"),
        ],
    },
    {
        "domain": "finance",
        "name": "finance_with_date",
        "title": "应收账款",
        "header_row": 2,
        "fields": [
            ("发票号", "invoice_no"),
            ("应收金额", "amount"),
            ("支付状态", "paid_status"),
            ("账户名称", "account_name"),
            ("业务日期", "order_date"),
        ],
    },
    {
        "domain": "finance",
        "name": "finance_multi_sheet",
        "title": "付款明细",
        "header_row": 2,
        "extra_sheet": True,
        "fields": [
            ("发票号", "invoice_no"),
            ("订单金额", "amount"),
            ("是否付款", "paid_status"),
            ("业务部门", "department"),
        ],
    },
]


def generate_validation_samples(output_dir: str | Path) -> list[ValidationSample]:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    return [_create_sample(target_dir, index, spec) for index, spec in enumerate(SPECS, start=1)]


def _create_sample(output_dir: Path, index: int, spec: dict) -> ValidationSample:
    path = output_dir / f"{index:02d}_{spec['name']}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec["domain"]
    fields = spec["fields"]
    header_row = spec["header_row"]
    if spec.get("title"):
        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(fields),
        )
        sheet.cell(row=1, column=1, value=spec["title"])
    if spec.get("section"):
        sheet.cell(row=header_row - 1, column=1, value=spec["section"])

    for column_index, (field_name, _target_field) in enumerate(fields, start=1):
        sheet.cell(row=header_row, column=column_index, value=field_name)
    for offset in range(3):
        row_index = header_row + 1 + offset
        for column_index, (_field_name, target_field) in enumerate(fields, start=1):
            sheet.cell(
                row=row_index,
                column=column_index,
                value=VALUE_BY_TARGET[target_field][offset],
            )
    for column_index in spec.get("hidden_columns", []):
        sheet.column_dimensions[get_column_letter(column_index)].hidden = True
    if spec.get("extra_sheet"):
        notes = workbook.create_sheet("说明")
        notes["A1"] = "此Sheet用于说明，不参与主模板验证"

    workbook.save(path)
    return ValidationSample(
        path=path,
        domain=spec["domain"],
        expected_header_row=header_row,
        expected_data_start_row=header_row + 1,
        expected_mappings={
            field_name: target_field
            for field_name, target_field in fields
            if target_field != "internal_note"
        },
    )
