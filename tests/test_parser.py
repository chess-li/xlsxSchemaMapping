from openpyxl import Workbook

from excel_matcher.excel.parser import parse_workbook


def test_parse_workbook_preserves_sheet_order_and_structural_cells(tmp_path):
    path = tmp_path / "structure.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "销售订单"
    sheet["A3"] = "订单号"
    sheet["B3"] = "客户名称"
    sheet["A4"] = "SO001"
    sheet["B4"] = "腾讯"
    sheet.column_dimensions["B"].hidden = True
    workbook.create_sheet("明细")
    workbook.save(path)

    parsed = parse_workbook(path)

    assert [sheet.name for sheet in parsed.sheets] == ["订单", "明细"]
    order_sheet = parsed.sheets[0]
    assert order_sheet.hidden_columns == ["B"]
    assert order_sheet.merged_ranges == ["A1:B1"]
    assert order_sheet.cells[0][0].coordinate == "A1"
    assert order_sheet.cells[0][0].display_value == "销售订单"
    assert order_sheet.cells[0][1].coordinate == "B1"
    assert order_sheet.cells[0][1].raw_value is None
    assert order_sheet.cells[0][1].display_value == "销售订单"
    assert order_sheet.cells[0][1].merged_anchor == "A1"
    assert [cell.display_value for cell in order_sheet.cells[1]] == [None, None]
