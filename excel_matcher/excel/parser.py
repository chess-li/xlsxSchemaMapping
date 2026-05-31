from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from excel_matcher.models import CellData, SheetData, WorkbookData


def parse_workbook(path: str | Path) -> WorkbookData:
    workbook_path = Path(path)
    workbook = load_workbook(workbook_path, data_only=True)
    sheets = [_parse_sheet(sheet) for sheet in workbook.worksheets]
    return WorkbookData(path=str(workbook_path), sheets=sheets)


def _parse_sheet(sheet) -> SheetData:
    merged_anchors = _merged_anchor_map(sheet)
    hidden_columns = [
        column_letter
        for column_letter, dimension in sheet.column_dimensions.items()
        if dimension.hidden
    ]
    cells: list[list[CellData]] = []
    for row_index in range(1, sheet.max_row + 1):
        row_cells = []
        for column_index in range(1, sheet.max_column + 1):
            coordinate = f"{get_column_letter(column_index)}{row_index}"
            cell = sheet.cell(row=row_index, column=column_index)
            merged_anchor = merged_anchors.get(coordinate)
            display_value = _display_value(sheet, cell.value, merged_anchor)
            row_cells.append(
                CellData(
                    row=row_index,
                    column=column_index,
                    coordinate=coordinate,
                    raw_value=cell.value,
                    display_value=display_value,
                    merged_anchor=merged_anchor,
                )
            )
        cells.append(row_cells)
    return SheetData(
        name=sheet.title,
        max_row=sheet.max_row,
        max_column=sheet.max_column,
        hidden_columns=hidden_columns,
        merged_ranges=[str(cell_range) for cell_range in sheet.merged_cells.ranges],
        cells=cells,
    )


def _merged_anchor_map(sheet) -> dict[str, str]:
    anchors: dict[str, str] = {}
    for merged_range in sheet.merged_cells.ranges:
        anchor = merged_range.start_cell.coordinate
        for row in sheet.iter_rows(
            min_row=merged_range.min_row,
            max_row=merged_range.max_row,
            min_col=merged_range.min_col,
            max_col=merged_range.max_col,
        ):
            for cell in row:
                if cell.coordinate != anchor:
                    anchors[cell.coordinate] = anchor
    return anchors


def _display_value(sheet, raw_value: Any, merged_anchor: str | None) -> Any:
    if merged_anchor is None:
        return raw_value
    return sheet[merged_anchor].value
